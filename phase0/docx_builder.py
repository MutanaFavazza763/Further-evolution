"""OCRDocument -> .docx 渲染器（Phase 0.6 最小实现）。

仅负责把结构化 OCRDocument 转成可编辑 Word，不涉及 GUI / 监控等模块。
数学公式走 LaTeX -> MathML -> OMML 链路，转成 Word 原生可编辑公式。

用法（供 verify_docx.py 调用）：
    render_document(data, output_path)
"""

import os
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree
from latex2mathml.converter import convert as latex_to_mathml

# MML2OMML.XSL 默认位置（本机 Office 自带），可用环境变量覆盖
DEFAULT_XSL_PATH = r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"

_transform_cache = {}


def _get_transform(xsl_path):
    if xsl_path not in _transform_cache:
        xslt_doc = etree.parse(xsl_path)
        _transform_cache[xsl_path] = etree.XSLT(xslt_doc)
    return _transform_cache[xsl_path]


def latex_to_omml(latex, xsl_path):
    """LaTeX -> OMML（lxml 元素）。失败抛出异常。"""
    mathml = latex_to_mathml(latex)
    transform = _get_transform(xsl_path)
    mathml_tree = etree.fromstring(mathml.encode("utf-8"))
    result = transform(mathml_tree)
    root = result.getroot()
    if root is None:
        raise ValueError("OMML 转换结果为空")
    return root


def _set_run_font(run, ascii_font="Times New Roman", east_asia_font="宋体"):
    """同时设置西文字体和中文字体。"""
    run.font.name = ascii_font
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), east_asia_font)


def _add_omml(paragraph, omml_element):
    paragraph._p.append(deepcopy(omml_element))


def _add_runs(paragraph, runs, xsl_path):
    """把 runs（text / formula 混排）按顺序渲染进一个段落。

    - 保持 text / formula 顺序
    - 文字与公式相邻时自动补空格，避免中英文紧贴
    - 纯公式（无相邻文字）不加空格，避免列表编号对齐问题
    """
    runs = runs or []
    for i, run in enumerate(runs):
        kind = run.get("kind")
        prev_is_text = i > 0 and runs[i - 1].get("kind") == "text"
        next_is_text = i < len(runs) - 1 and runs[i + 1].get("kind") == "text"

        if kind == "text":
            text = run.get("text", "")
            if text:
                r = paragraph.add_run(text)
                _set_run_font(r)
        elif kind == "formula":
            latex = run.get("latex", "")
            if not latex:
                continue
            try:
                omml = latex_to_omml(latex, xsl_path)
                prev_text = runs[i - 1].get("text", "") if prev_is_text else ""
                next_text = runs[i + 1].get("text", "") if next_is_text else ""
                # 仅在相邻文字本身不含首尾空格时补一个空格，避免双空格
                if prev_text and not prev_text.endswith((" ", "\u3000")):
                    paragraph.add_run(" ")
                _add_omml(paragraph, omml)
                if next_text and not next_text.startswith((" ", "\u3000")):
                    r = paragraph.add_run(" ")
                    _set_run_font(r)
            except Exception:
                # 转换失败降级为纯文本，不中断整个文档
                r = paragraph.add_run(" " + latex + " ")
                _set_run_font(r)


def _add_table(doc, rows, xsl_path):
    n_rows = len(rows)
    n_cols = max((len(r) for r in rows), default=0)
    if n_cols == 0:
        return
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(n_cols):
            cell_data = row[j] if j < len(row) else None
            p = table.cell(i, j).paragraphs[0]
            if isinstance(cell_data, str):
                r = p.add_run(cell_data)
                _set_run_font(r)
            elif isinstance(cell_data, dict):
                _add_runs(p, cell_data.get("runs"), xsl_path)


def render_document(data, output_path, xsl_path=None):
    xsl_path = xsl_path or os.environ.get(
        "OCR2WORD_MML2OMML_XSL", DEFAULT_XSL_PATH
    )
    doc = Document()

    for block in data.get("blocks", []):
        btype = block.get("type")

        if btype == "heading":
            text = block.get("text", "")
            if not text:
                continue
            level = max(1, min(6, int(block.get("level", 1))))
            h = doc.add_heading(level=level)
            r = h.add_run(text)
            _set_run_font(r)

        elif btype in ("paragraph", "handwritten"):
            runs = block.get("runs") or []
            if not runs:
                continue
            p = doc.add_paragraph()
            _add_runs(p, runs, xsl_path)

        elif btype == "list":
            ordered = bool(block.get("ordered", False))
            style_name = "List Number" if ordered else "List Bullet"
            for item in block.get("items", []):
                p = doc.add_paragraph(style=style_name)
                if isinstance(item, str):
                    r = p.add_run(item)
                    _set_run_font(r)
                else:
                    _add_runs(p, item, xsl_path)

        elif btype == "formula":
            latex = block.get("latex", "")
            if not latex:
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                omml = latex_to_omml(latex, xsl_path)
                _add_omml(p, omml)
            except Exception:
                r = p.add_run(latex)
                _set_run_font(r)

        elif btype == "table":
            _add_table(doc, block.get("rows", []), xsl_path)

        elif btype == "image":
            # 第一版跳过：schema 缺 bbox，无法裁剪原图区域
            p = doc.add_paragraph()
            r = p.add_run("[图片块：暂未实现]")
            _set_run_font(r)

    doc.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# Phase 2：横向 16:9 重排渲染
# ---------------------------------------------------------------------------

def _setup_landscape_16_9(doc):
    """设置横向 + 16:9 页面尺寸（13.333 × 7.5 英寸）。"""
    section = doc.sections[0]
    section.orientation = WD_ORIENTATION.LANDSCAPE
    section.page_width = Inches(13.333)
    section.page_height = Inches(7.5)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)


def _render_block(container, block, xsl_path):
    """把单个 block 渲染到容器（Document 或 _Cell）。"""
    btype = block.get("type")
    if btype == "heading":
        text = block.get("text", "")
        if not text:
            return
        level = max(1, min(6, int(block.get("level", 1))))
        p = container.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt({1: 22, 2: 18, 3: 15}.get(level, 15))
        _set_run_font(r)

    elif btype in ("paragraph", "handwritten"):
        runs = block.get("runs") or []
        if not runs:
            return
        p = container.add_paragraph()
        _add_runs(p, runs, xsl_path)

    elif btype == "list":
        for item in block.get("items", []):
            p = container.add_paragraph()
            _add_runs(p, item, xsl_path)

    elif btype == "formula":
        latex = block.get("latex", "")
        if not latex:
            return
        p = container.add_paragraph()
        try:
            omml = latex_to_omml(latex, xsl_path)
            _add_omml(p, omml)
        except Exception:
            r = p.add_run(latex)
            _set_run_font(r)

    elif btype == "image":
        p = container.add_paragraph()
        r = p.add_run("[图片]")
        _set_run_font(r)

    elif btype == "table":
        # 嵌套表格：MVP 简化为文本行
        for row in block.get("rows", []):
            parts = []
            for cell in row:
                if isinstance(cell, str):
                    parts.append(cell)
                elif isinstance(cell, dict):
                    parts.append("".join(
                        r.get("text", "") or r.get("latex", "")
                        for r in cell.get("runs", [])
                    ))
            p = container.add_paragraph()
            r = p.add_run(" | ".join(parts))
            _set_run_font(r)


def render_reflow(layout, columns, output_path, xsl_path=None):
    """把版面分析结果渲染成横向 16:9 多栏 DOCX。

    layout：phase2.layout.PageLayout
    columns：list[list[BlockGroup]]，来自 phase2.reflow.reflow
    """
    xsl_path = xsl_path or os.environ.get(
        "OCR2WORD_MML2OMML_XSL", DEFAULT_XSL_PATH
    )
    doc = Document()
    _setup_landscape_16_9(doc)

    # header（顶部，跨栏）
    for block in layout.header:
        _render_block(doc, block, xsl_path)

    # main：表格模拟多栏
    n_cols = len(columns)
    if n_cols > 0:
        table = doc.add_table(rows=1, cols=n_cols)
        for col_idx, col_groups in enumerate(columns):
            cell = table.cell(0, col_idx)
            first_p = cell.paragraphs[0]._p
            first_p.getparent().remove(first_p)
            for group in col_groups:
                for block in group.blocks:
                    _render_block(cell, block, xsl_path)

    # footer（底部，跨栏）
    for block in layout.footer:
        _render_block(doc, block, xsl_path)

    doc.save(output_path)
    return output_path
