"""Phase 3 端到端测试：bbox 完整链路 order_blocks → layout → reflow → DOCX。

使用 double_column_bbox.json，不绕过任何实际 pipeline / renderer 逻辑。
验证：标题渲染为正文段落（表格外），左右栏内容渲染进多栏表格单元格。
"""

import json
import sys
import tempfile
import unittest
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(_ROOT))
sys.path.insert(0, str(_ROOT / "phase0"))

from docx_builder import render_reflow
from phase2.layout import analyze_layout
from phase2.ordering import order_blocks, order_groups
from phase2.reflow import choose_columns, reflow

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def _load(name):
    return json.loads((FIXTURES / name).read_text(encoding="utf-8"))


def _label(block):
    t = block.get("type")
    if t == "heading":
        return block.get("text", "")
    if t in ("paragraph", "handwritten"):
        runs = block.get("runs", [])
        if runs:
            return runs[0].get("text", "") or runs[0].get("latex", "")
    return ""


class TestPipelineBbox(unittest.TestCase):
    def test_full_bbox_pipeline(self):
        blocks = _load("double_column_bbox.json")["blocks"]

        # 1. 阅读顺序重建
        ordered = order_blocks(blocks)
        self.assertEqual(
            [_label(b) for b in ordered],
            ["标题", "1", "2", "3", "4", "5"],
        )

        # 2. 版面分析：spanning 含标题，main 不含标题
        layout = analyze_layout(ordered)
        self.assertEqual(
            [_label(b) for g in layout.spanning for b in g.blocks],
            ["标题"],
        )
        self.assertEqual(
            [_label(b) for g in layout.main for b in g.blocks],
            ["1", "2", "3", "4", "5"],
        )

        # 3. reflow：main 内容不丢失
        groups = order_groups(layout.main)
        columns = reflow(groups, columns=choose_columns(len(groups)))
        col_blocks = [b for col in columns for g in col for b in g.blocks]
        self.assertEqual(len(col_blocks), 5)

        # 4. 生成 DOCX
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "bbox_test.docx"
            render_reflow(layout, columns, str(output))

            from docx import Document
            doc = Document(str(output))

            # 5. 区分正文段落与表格单元格
            body_texts = [p.text for p in doc.paragraphs]
            self.assertIn("标题", "".join(body_texts))

            cell_texts = []
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            cell_texts.append(p.text)
            cell_joined = "".join(cell_texts)

            # 标题不在表格单元格里（是正文段落）
            self.assertNotIn("标题", cell_joined)

            # 左右栏内容都在表格单元格里
            for num in ["1", "2", "3", "4", "5"]:
                self.assertIn(num, cell_joined)


if __name__ == "__main__":
    unittest.main()
